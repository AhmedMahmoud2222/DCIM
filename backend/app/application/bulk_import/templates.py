"""Downloadable XLSX templates for the three bulk-import types, built with openpyxl.
Each template is a header row (bold, frozen), one example data row, an "Instructions"
sheet documenting every column, and a data-validation dropdown for every enum column —
openpyxl's `DataValidation` — since Excel/LibreOffice enforce those client-side, well
before the row ever reaches `parsing.py`/the validators.

`RACK_COLUMNS`/`EQUIPMENT_COLUMNS`/`CATALOG_COLUMNS` are the single source of truth for
column name, order, and required-ness — `parsing.py` (header validation),
`report.py` (rebuilding the original columns), and the three validators all import from
here rather than repeating the column list, so the template, the parser, and the report
can never silently drift apart."""

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation


@dataclass(frozen=True)
class ColumnSpec:
    name: str
    required: bool
    description: str
    choices: tuple[str, ...] | None = None


RACK_COLUMNS: tuple[ColumnSpec, ...] = (
    ColumnSpec("asset_tag", True, "Globally unique asset tag for the rack."),
    ColumnSpec("rack_name", True, "Display name for the rack."),
    ColumnSpec("manufacturer", True, "Manufacturer name — must match an existing catalog manufacturer."),
    ColumnSpec("model_name", True, "Rack model name — must match an existing published rack model revision."),
    ColumnSpec(
        "revision_number", False,
        "Legacy rack model revision number. Leave blank to use the latest available revision.",
    ),
    ColumnSpec("site_code", True, "Site code (Site.code)."),
    ColumnSpec("building_code", True, "Building code, unique within the site."),
    ColumnSpec("floor_level", True, "Floor level number, unique within the building."),
    ColumnSpec("room_code", True, "Room code, unique within the floor."),
    ColumnSpec("x_mm", False, "Placement X coordinate in mm, if known."),
    ColumnSpec("y_mm", False, "Placement Y coordinate in mm, if known."),
    ColumnSpec("rotation_deg", False, "Placement rotation in degrees [0, 360)."),
    ColumnSpec("owner", False, "Free-text owner/team."),
    ColumnSpec("notes", False, "Free-text notes."),
)

EQUIPMENT_COLUMNS: tuple[ColumnSpec, ...] = (
    ColumnSpec("asset_tag", True, "Globally unique asset tag for the equipment."),
    ColumnSpec("hostname", False, "Hostname, if any."),
    ColumnSpec("manufacturer", True, "Manufacturer name — must match an existing catalog manufacturer."),
    ColumnSpec("model_name", True, "Equipment model name — must match an existing published equipment model revision."),
    ColumnSpec(
        "revision_number", False,
        "Legacy equipment model revision number. Leave blank to use the latest available revision.",
    ),
    ColumnSpec(
        "placement_type", True, "How the equipment is mounted.",
        choices=("rack_mounted", "floor_standing", "wall_mounted", "ceiling_mounted", "other"),
    ),
    ColumnSpec("site_code", True, "Site code (Site.code)."),
    ColumnSpec("building_code", True, "Building code, unique within the site."),
    ColumnSpec("floor_level", True, "Floor level number, unique within the building."),
    ColumnSpec("room_code", True, "Room code, unique within the floor."),
    ColumnSpec("rack_asset_tag", False, "Asset tag of the target rack. Required when placement_type=rack_mounted."),
    ColumnSpec("u_start", False, "Starting U position (1-based). Required when placement_type=rack_mounted."),
    ColumnSpec("u_end", False, "Ending U position (exclusive). Required when placement_type=rack_mounted."),
    ColumnSpec(
        "side", False, "Rack side occupied. Required when placement_type=rack_mounted.", choices=("front", "rear", "both"),
    ),
    ColumnSpec("ip_address", False, "Management IP address, if any."),
    ColumnSpec("mac_address", False, "MAC address, if any."),
    ColumnSpec("owner", False, "Free-text owner/team."),
    ColumnSpec("service", False, "Free-text service name."),
    ColumnSpec("environment", False, "Free-text environment (e.g. production, staging)."),
    ColumnSpec("notes", False, "Free-text notes."),
    ColumnSpec(
        "lifecycle_status", False, "Initial lifecycle status. Defaults to 'planned' when left blank.",
        choices=("planned", "installed", "active", "reserved", "maintenance", "decommissioned", "removed"),
    ),
)

CATALOG_COLUMNS: tuple[ColumnSpec, ...] = (
    ColumnSpec("manufacturer_name", True, "Manufacturer name. Created if it does not already exist."),
    ColumnSpec(
        "category", True, "Catalog category.",
        choices=("rack", "equipment", "network_device", "pdu", "ups", "power_panel", "sensor"),
    ),
    ColumnSpec("model_name", True, "Catalog model name."),
    ColumnSpec("model_number", False, "Manufacturer model/part number."),
    ColumnSpec("subtype", False, "Free-text subtype."),
    ColumnSpec("description", False, "Free-text description."),
    ColumnSpec("tags", False, "Comma-separated tags."),
    ColumnSpec("dimension_unit", False, "Unit for width/height/depth.", choices=("mm", "in")),
    ColumnSpec("width_value", False, "Width, in dimension_unit."),
    ColumnSpec("height_value", False, "Height, in dimension_unit."),
    ColumnSpec("depth_value", False, "Depth, in dimension_unit."),
    ColumnSpec("rack_unit_height", False, "Rack unit (U) height, for rack/rack-mountable equipment models."),
    ColumnSpec("weight_unit", False, "Unit for weight_value.", choices=("kg", "lb")),
    ColumnSpec("weight_value", False, "Weight, in weight_unit."),
    ColumnSpec("mounting_orientation", False, "Free-text mounting orientation."),
    ColumnSpec(
        "supported_placement_types", False,
        "Comma-separated placement types (rack_mounted, floor_standing, wall_mounted, ceiling_mounted, other).",
    ),
    ColumnSpec(
        "airflow_direction", False, "Airflow direction.",
        choices=("front_to_rear", "front_to_top", "side_to_side", "other"),
    ),
    ColumnSpec("rated_power_w", False, "Rated power, in watts."),
    ColumnSpec("typical_power_w", False, "Typical power, in watts."),
    ColumnSpec("max_power_w", False, "Maximum power, in watts."),
    ColumnSpec("heat_dissipation_btu_hr", False, "Heat dissipation, in BTU/hr."),
    ColumnSpec("power_redundancy_mode", False, "Power redundancy mode.", choices=("single", "1+1", "n+1")),
    ColumnSpec(
        "revision_number", False,
        "Only meaningful in update_existing mode: the draft revision number to update.",
    ),
    ColumnSpec(
        "clone_from_revision_number", False,
        "Only meaningful in create_only mode: seed a new draft by cloning this existing published/retired revision.",
    ),
)

EXAMPLE_ROWS: dict[str, dict[str, object]] = {
    "rack": {
        "asset_tag": "RACK-0001", "rack_name": "Row A Rack 1", "manufacturer": "Acme", "model_name": "RM-42U",
        "revision_number": "", "site_code": "SITE1", "building_code": "A", "floor_level": 1, "room_code": "R101",
        "x_mm": 0, "y_mm": 0, "rotation_deg": 0, "owner": "Facilities", "notes": "",
    },
    "equipment": {
        "asset_tag": "EQ-0001", "hostname": "srv-0001", "manufacturer": "Acme", "model_name": "EM-1U",
        "revision_number": "", "placement_type": "rack_mounted", "site_code": "SITE1", "building_code": "A",
        "floor_level": 1, "room_code": "R101", "rack_asset_tag": "RACK-0001", "u_start": 1, "u_end": 2,
        "side": "front", "ip_address": "", "mac_address": "", "owner": "", "service": "", "environment": "",
        "notes": "", "lifecycle_status": "planned",
    },
    "catalog": {
        "manufacturer_name": "Acme", "category": "equipment", "model_name": "EM-1U", "model_number": "EM-1U-A1",
        "subtype": "", "description": "", "tags": "", "dimension_unit": "mm", "width_value": 440, "height_value": 44,
        "depth_value": 600, "rack_unit_height": 1, "weight_unit": "kg", "weight_value": 5, "mounting_orientation": "",
        "supported_placement_types": "rack_mounted", "airflow_direction": "front_to_rear", "rated_power_w": 200,
        "typical_power_w": 120, "max_power_w": 250, "heat_dissipation_btu_hr": "", "power_redundancy_mode": "single",
        "revision_number": "", "clone_from_revision_number": "",
    },
}


def _build_template(*, sheet_title: str, columns: tuple[ColumnSpec, ...], example: dict[str, object]) -> bytes:
    workbook = Workbook()
    data_sheet = workbook.active
    assert data_sheet is not None
    data_sheet.title = sheet_title

    header_font = Font(bold=True)
    for col_index, column in enumerate(columns, start=1):
        cell = data_sheet.cell(row=1, column=col_index, value=column.name)
        cell.font = header_font
    data_sheet.freeze_panes = "A2"

    for col_index, column in enumerate(columns, start=1):
        data_sheet.cell(row=2, column=col_index, value=example.get(column.name, ""))

    for col_index, column in enumerate(columns, start=1):
        if column.choices:
            validation = DataValidation(
                type="list", formula1=f'"{",".join(column.choices)}"', allow_blank=not column.required,
            )
            data_sheet.add_data_validation(validation)
            col_letter = data_sheet.cell(row=1, column=col_index).column_letter
            validation.add(f"{col_letter}2:{col_letter}1048576")

    instructions = workbook.create_sheet("Instructions")
    instructions.cell(row=1, column=1, value="Column").font = header_font
    instructions.cell(row=1, column=2, value="Required").font = header_font
    instructions.cell(row=1, column=3, value="Description").font = header_font
    instructions.cell(row=1, column=4, value="Valid values").font = header_font
    for row_index, column in enumerate(columns, start=2):
        instructions.cell(row=row_index, column=1, value=column.name)
        instructions.cell(row=row_index, column=2, value="yes" if column.required else "no")
        instructions.cell(row=row_index, column=3, value=column.description)
        instructions.cell(row=row_index, column=4, value=", ".join(column.choices) if column.choices else "")
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["C"].width = 70
    instructions.column_dimensions["D"].width = 40

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_rack_template() -> bytes:
    return _build_template(sheet_title="Racks", columns=RACK_COLUMNS, example=EXAMPLE_ROWS["rack"])


def build_equipment_template() -> bytes:
    return _build_template(sheet_title="Equipment", columns=EQUIPMENT_COLUMNS, example=EXAMPLE_ROWS["equipment"])


def build_catalog_template() -> bytes:
    return _build_template(sheet_title="Catalog", columns=CATALOG_COLUMNS, example=EXAMPLE_ROWS["catalog"])


COLUMNS_BY_IMPORT_TYPE: dict[str, tuple[ColumnSpec, ...]] = {
    "rack": RACK_COLUMNS,
    "equipment": EQUIPMENT_COLUMNS,
    "catalog": CATALOG_COLUMNS,
}
