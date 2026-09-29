"""Equipment bulk-import pipeline — see test_bulk_import_racks.py's module docstring for
the `.run()` synchronous-Celery-task pattern this file also uses."""

import uuid

from app.infrastructure.tasks.bulk_import import commit_bulk_import_job, parse_and_validate_bulk_import_job
from tests.api._bulk_import_helpers import build_workbook, create_room_with_codes
from tests.api._phase2_helpers import create_equipment, create_rack

EQUIPMENT_HEADERS = [
    "asset_tag", "hostname", "manufacturer", "model_name", "revision_number", "placement_type", "site_code",
    "building_code", "floor_level", "room_code", "rack_asset_tag", "u_start", "u_end", "side", "ip_address",
    "mac_address", "owner", "service", "environment", "notes", "lifecycle_status",
]


async def _create_equipment_model(client, auth_headers) -> dict:
    admin_headers = await auth_headers("Administrator")
    model_name = f"EM-{uuid.uuid4().hex[:8]}"
    model = (
        await client.post(
            "/api/v1/equipment-models", json={"manufacturer": "Acme", "model_name": model_name}, headers=admin_headers,
        )
    ).json()
    revision = (
        await client.post(f"/api/v1/equipment-models/{model['id']}/revisions", json={}, headers=admin_headers)
    ).json()
    return {"manufacturer": "Acme", "model_name": model_name, "revision_id": revision["id"]}


def _run_parse(job_id: str) -> None:
    parse_and_validate_bulk_import_job.run(job_id)


async def _run_commit(client, headers, job_id: str) -> None:
    """See test_bulk_import_racks.py's identical helper for the full reasoning (SEC,
    Codex PR #50 review, finding #1)."""
    resp = await client.post(f"/api/v1/import-jobs/{job_id}/commit", headers=headers)
    assert resp.status_code == 202, resp.text
    commit_bulk_import_job.run(job_id)


def _floor_standing_row(asset_tag, model, room) -> list:
    return [
        asset_tag, "host-1", model["manufacturer"], model["model_name"], "", "floor_standing", room["site_code"],
        room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "", "", "", "", "", "", "planned",
    ]


async def test_equipment_import_template_download_is_real_xlsx_with_expected_headers(client, auth_headers):
    headers = await auth_headers("Engineer")
    resp = await client.get("/api/v1/equipment/import-template", headers=headers)
    assert resp.status_code == 200
    assert resp.content[:4] == b"PK\x03\x04"

    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(resp.content))
    sheet = workbook["Equipment"]
    header_row = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header_row == EQUIPMENT_HEADERS


async def test_upload_parse_and_preview_rows(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    asset_tag = f"EQ-{uuid.uuid4().hex[:8]}"

    content = build_workbook(EQUIPMENT_HEADERS, [_floor_standing_row(asset_tag, model, room)])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=create_only",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    assert upload.status_code == 202, upload.text
    job = upload.json()

    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "validated"
    assert job_status["valid_row_count"] == 1


async def test_commit_create_only_mode_success_rack_mounted(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    rack = await create_rack(client, headers, auth_headers, room_id=room["room_id"])
    asset_tag = f"EQ-{uuid.uuid4().hex[:8]}"

    row = [
        asset_tag, "srv-1", model["manufacturer"], model["model_name"], "", "rack_mounted", room["site_code"],
        room["building_code"], room["floor_level"], room["room_code"], rack["asset_tag"], 1, 2, "front", "", "",
        "Ops", "", "", "", "planned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=create_only",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    await _run_commit(client, headers, job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status
    assert job_status["committed_row_count"] == 1

    items = (await client.get("/api/v1/equipment", headers=headers)).json()["items"]
    created = next(e for e in items if e["asset_tag"] == asset_tag)
    assert created["placement"]["placement_type"] == "rack_mounted"
    assert created["placement"]["rack_id"] == rack["id"]
    assert created["placement"]["u_start"] == 1
    assert created["placement"]["u_end"] == 2


async def test_commit_update_existing_mode_success(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    existing = await create_equipment(client, headers, auth_headers)

    row = [
        existing["asset_tag"], "new-hostname", model["manufacturer"], model["model_name"], "", "floor_standing",
        room["site_code"], room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "", "",
        "New Owner", "", "", "updated", "planned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=update_existing",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    await _run_commit(client, headers, job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status

    updated = (await client.get(f"/api/v1/equipment/{existing['id']}", headers=headers)).json()
    assert updated["hostname"] == "new-hostname"
    assert updated["owner"] == "New Owner"
    assert updated["model_revision_id"] == existing["model_revision_id"]


async def test_update_existing_row_fails_if_equipment_changed_since_preview(client, auth_headers):
    """SEC (Codex PR #50 review, finding #5): see test_bulk_import_racks.py's identical
    test — same reasoning, applied to Equipment.version."""
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    existing = await create_equipment(client, headers, auth_headers)
    assert existing["version"] == 1

    row = [
        existing["asset_tag"], "new-hostname", model["manufacturer"], model["model_name"], "", "floor_standing",
        room["site_code"], room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "", "",
        "Stale Owner", "", "", "", "planned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=update_existing",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    patch = await client.patch(
        f"/api/v1/equipment/{existing['id']}", json={"owner": "Concurrent Edit"},
        headers={**headers, "If-Match": "1"},
    )
    assert patch.status_code == 200, patch.text
    assert patch.json()["version"] == 2

    await _run_commit(client, headers, job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["committed_row_count"] == 0
    assert job_status["failed_row_count"] == 1

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    failed_row = rows["items"][0]
    assert failed_row["status"] == "failed"
    assert "changed" in failed_row["errors"][-1]["message"].lower()

    current = (await client.get(f"/api/v1/equipment/{existing['id']}", headers=headers)).json()
    assert current["owner"] == "Concurrent Edit"
    assert current["version"] == 2


async def test_lifecycle_status_change_on_update_is_rejected_not_silently_dropped(client, auth_headers):
    """SEC (Codex PR #50 review, finding #6): lifecycle_status has no update path
    anywhere else in this application — a bulk-update row attempting to change it must
    be a clear row-level rejection, never a silent no-op."""
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    existing = await create_equipment(client, headers, auth_headers)  # lifecycle_status defaults to "planned"

    row = [
        existing["asset_tag"], "new-hostname", model["manufacturer"], model["model_name"], "", "floor_standing",
        room["site_code"], room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "", "",
        "", "", "", "", "decommissioned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=update_existing",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    await _run_commit(client, headers, job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["committed_row_count"] == 0
    assert job_status["failed_row_count"] == 1

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert "lifecycle_status" in rows["items"][0]["errors"][-1]["message"].lower()

    current = (await client.get(f"/api/v1/equipment/{existing['id']}", headers=headers)).json()
    assert current["lifecycle_status"] == "planned"


async def test_update_existing_applies_mac_address_and_ip_address(client, auth_headers, db_session):
    """SEC (Codex PR #50 review, finding #6): mac_address/ip_address were previously
    silently dropped by commit (mac_address on both create and update; ip_address on
    update only). EquipmentOut does not expose either column (a pre-existing,
    out-of-scope gap), so this reads the committed row directly."""
    from sqlalchemy import text as sql_text

    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    asset_tag = f"EQ-{uuid.uuid4().hex[:8]}"

    create_row = [
        asset_tag, "host-1", model["manufacturer"], model["model_name"], "", "floor_standing", room["site_code"],
        room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "10.0.0.1", "AA:BB:CC:DD:EE:01",
        "", "", "", "", "planned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [create_row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=create_only",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    await _run_commit(client, headers, job["id"])

    items = (await client.get("/api/v1/equipment", headers=headers)).json()["items"]
    created = next(e for e in items if e["asset_tag"] == asset_tag)

    async def _fetch(equipment_id: str) -> dict:
        row = (
            await db_session.execute(
                sql_text("SELECT ip_address, mac_address FROM equipment WHERE id = :id"), {"id": equipment_id}
            )
        ).mappings().one()
        # ip_address is INET at the DB level — asyncpg/SQLAlchemy surfaces it as an
        # ipaddress.IPv4Address, not a plain str.
        return {"ip_address": str(row["ip_address"]), "mac_address": row["mac_address"]}

    after_create = await _fetch(created["id"])
    assert after_create["ip_address"] == "10.0.0.1"
    assert after_create["mac_address"] == "AA:BB:CC:DD:EE:01"

    update_row = [
        asset_tag, "host-1", model["manufacturer"], model["model_name"], "", "floor_standing", room["site_code"],
        room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", "10.0.0.2", "AA:BB:CC:DD:EE:02",
        "", "", "", "", "planned",
    ]
    update_content = build_workbook(EQUIPMENT_HEADERS, [update_row])
    update_upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=update_existing",
        files={"file": ("equipment.xlsx", update_content, "application/octet-stream")}, headers=headers,
    )
    update_job = update_upload.json()
    _run_parse(update_job["id"])
    await _run_commit(client, headers, update_job["id"])

    after_update = await _fetch(created["id"])
    assert after_update["ip_address"] == "10.0.0.2"
    assert after_update["mac_address"] == "AA:BB:CC:DD:EE:02"


async def test_invalid_rack_asset_tag(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)

    row = [
        f"EQ-{uuid.uuid4().hex[:8]}", "srv-1", model["manufacturer"], model["model_name"], "", "rack_mounted",
        room["site_code"], room["building_code"], room["floor_level"], room["room_code"], "NO-SUCH-RACK", 1, 2,
        "front", "", "", "", "", "", "", "planned",
    ]
    content = build_workbook(EQUIPMENT_HEADERS, [row])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=create_only",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert rows["items"][0]["status"] == "invalid"
    assert rows["items"][0]["errors"][0]["field"] == "rack_asset_tag"


async def test_overlapping_u_positions_between_two_rows_targeting_the_same_rack(client, auth_headers):
    """Overlap is not a validate-time rejection (each row is individually plausible) —
    it is caught at commit time by the database's GiST exclusion constraint on
    equipment_placement: exactly one of the two rows commits, the other fails."""
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_equipment_model(client, auth_headers)
    rack = await create_rack(client, headers, auth_headers, room_id=room["room_id"])

    def _row(tag):
        return [
            tag, tag, model["manufacturer"], model["model_name"], "", "rack_mounted", room["site_code"],
            room["building_code"], room["floor_level"], room["room_code"], rack["asset_tag"], 1, 2, "front", "", "",
            "", "", "", "", "planned",
        ]

    tag_a, tag_b = f"EQ-{uuid.uuid4().hex[:8]}", f"EQ-{uuid.uuid4().hex[:8]}"
    content = build_workbook(EQUIPMENT_HEADERS, [_row(tag_a), _row(tag_b)])
    upload = await client.post(
        "/api/v1/equipment/import-jobs?mode=create_only",
        files={"file": ("equipment.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    await _run_commit(client, headers, job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed_with_errors", job_status
    assert job_status["committed_row_count"] == 1
    assert job_status["failed_row_count"] == 1

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    statuses = {r["raw_data"]["asset_tag"]: r["status"] for r in rows["items"]}
    assert sorted(statuses.values()) == ["committed", "failed"]
    failed_row = next(r for r in rows["items"] if r["status"] == "failed")
    assert "overlap" in failed_row["errors"][-1]["message"].lower()
