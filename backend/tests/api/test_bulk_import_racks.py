"""Rack bulk-import pipeline (backend plan: shared bulk-import job pipeline, rack
plug-in). Celery tasks are invoked synchronously via `.run()` — the same pattern
tests/api/test_floor_plans.py already established for this test environment (no worker
consumes the broker queue here)."""

import uuid

from app.infrastructure.tasks.bulk_import import commit_bulk_import_job, parse_and_validate_bulk_import_job
from tests.api._bulk_import_helpers import build_workbook, create_room_with_codes
from tests.api._phase2_helpers import create_rack

RACK_HEADERS = [
    "asset_tag", "rack_name", "manufacturer", "model_name", "revision_number", "site_code", "building_code",
    "floor_level", "room_code", "x_mm", "y_mm", "rotation_deg", "owner", "notes",
]


async def _create_rack_model(client, auth_headers) -> dict:
    admin_headers = await auth_headers("Administrator")
    model_name = f"RM-{uuid.uuid4().hex[:8]}"
    model = (
        await client.post(
            "/api/v1/rack-models", json={"manufacturer": "Acme", "model_name": model_name}, headers=admin_headers,
        )
    ).json()
    revision = (
        await client.post(
            f"/api/v1/rack-models/{model['id']}/revisions",
            json={"height_u": 42, "width_mm": 600, "depth_mm": 1000}, headers=admin_headers,
        )
    ).json()
    return {"manufacturer": "Acme", "model_name": model_name, "revision_id": revision["id"]}


def _run_parse(job_id: str) -> None:
    parse_and_validate_bulk_import_job.run(job_id)


def _run_commit(job_id: str) -> None:
    commit_bulk_import_job.run(job_id)


async def test_rack_import_template_download_is_real_xlsx_with_expected_headers(client, auth_headers):
    headers = await auth_headers("Engineer")
    resp = await client.get("/api/v1/racks/import-template", headers=headers)
    assert resp.status_code == 200
    assert resp.content[:4] == b"PK\x03\x04"

    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(resp.content))
    sheet = workbook["Racks"]
    header_row = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header_row == RACK_HEADERS
    assert "Instructions" in workbook.sheetnames


async def test_upload_parse_and_preview_rows(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    asset_tag = f"RACK-{uuid.uuid4().hex[:8]}"

    content = build_workbook(
        RACK_HEADERS,
        [[asset_tag, "Row A Rack 1", model["manufacturer"], model["model_name"], "", room["site_code"],
          room["building_code"], room["floor_level"], room["room_code"], 0, 0, 0, "Facilities", ""]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    assert upload.status_code == 202, upload.text
    job = upload.json()
    assert job["status"] == "uploaded"

    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "validated", job_status
    assert job_status["row_count"] == 1
    assert job_status["valid_row_count"] == 1
    assert job_status["error_row_count"] == 0

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert rows["total"] == 1
    assert rows["items"][0]["status"] == "valid"
    assert rows["items"][0]["raw_data"]["asset_tag"] == asset_tag


async def test_commit_create_only_mode_success(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    asset_tag = f"RACK-{uuid.uuid4().hex[:8]}"

    content = build_workbook(
        RACK_HEADERS,
        [[asset_tag, "Row A Rack 1", model["manufacturer"], model["model_name"], "", room["site_code"],
          room["building_code"], room["floor_level"], room["room_code"], 100, 200, 0, "Facilities", "note"]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    _run_commit(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status
    assert job_status["committed_row_count"] == 1
    assert job_status["failed_row_count"] == 0
    assert job_status["report_available"] is True

    racks = (await client.get("/api/v1/racks", headers=headers)).json()
    created = next(r for r in racks["items"] if r["asset_tag"] == asset_tag)
    assert created["name"] == "Row A Rack 1"
    assert created["owner"] == "Facilities"
    assert created["placement"]["room_id"] == room["room_id"]
    assert created["placement"]["x_mm"] == 100

    report = await client.get(f"/api/v1/import-jobs/{job['id']}/report", headers=headers)
    assert report.status_code == 200
    assert report.content[:4] == b"PK\x03\x04"


async def test_commit_update_existing_mode_success(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    existing = await create_rack(client, headers, auth_headers, room_id=room["room_id"])

    content = build_workbook(
        RACK_HEADERS,
        [[existing["asset_tag"], existing["name"], model["manufacturer"], model["model_name"], "", room["site_code"],
          room["building_code"], room["floor_level"], room["room_code"], 5, 5, 0, "New Owner", "updated"]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=update_existing",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    _run_commit(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status
    assert job_status["committed_row_count"] == 1

    updated = (await client.get(f"/api/v1/racks/{existing['id']}", headers=headers)).json()
    assert updated["owner"] == "New Owner"
    assert updated["notes"] == "updated"
    # model_revision_id must never be silently re-homed by a bulk update.
    assert updated["model_revision_id"] == existing["model_revision_id"]


async def test_duplicate_asset_tag_within_the_same_file(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    asset_tag = f"RACK-{uuid.uuid4().hex[:8]}"
    row = [asset_tag, "Rack", model["manufacturer"], model["model_name"], "", room["site_code"], room["building_code"],
           room["floor_level"], room["room_code"], "", "", "", "", ""]

    content = build_workbook(RACK_HEADERS, [row, row])
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["row_count"] == 2
    assert job_status["valid_row_count"] == 1
    assert job_status["error_row_count"] == 1

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows?status=invalid", headers=headers)).json()
    assert rows["total"] == 1
    assert "duplicate" in rows["items"][0]["errors"][0]["message"].lower()


async def test_duplicate_asset_tag_against_existing_db_row(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    existing = await create_rack(client, headers, auth_headers, room_id=room["room_id"])

    content = build_workbook(
        RACK_HEADERS,
        [[existing["asset_tag"], "Another Rack", model["manufacturer"], model["model_name"], "", room["site_code"],
          room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", ""]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["error_row_count"] == 1
    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert "already exists" in rows["items"][0]["errors"][0]["message"].lower()


async def test_invalid_catalog_model_reference(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)

    content = build_workbook(
        RACK_HEADERS,
        [[f"RACK-{uuid.uuid4().hex[:8]}", "Rack", "NoSuchManufacturer", "NoSuchModel", "", room["site_code"],
          room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", ""]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["error_row_count"] == 1
    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert "no such rack model" in rows["items"][0]["errors"][0]["message"].lower()


async def test_nonexistent_room_path(client, auth_headers):
    headers = await auth_headers("Engineer")
    model = await _create_rack_model(client, auth_headers)

    content = build_workbook(
        RACK_HEADERS,
        [[f"RACK-{uuid.uuid4().hex[:8]}", "Rack", model["manufacturer"], model["model_name"], "", "NOSITE", "NOBLDG",
          1, "NOROOM", "", "", "", "", ""]],
    )
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    rows = (await client.get(f"/api/v1/import-jobs/{job['id']}/rows", headers=headers)).json()
    assert rows["items"][0]["status"] == "invalid"
    assert rows["items"][0]["errors"][0]["field"] == "site_code"


async def test_oversized_file_rejected_413(client, auth_headers):
    headers = await auth_headers("Engineer")
    oversized = b"PK\x03\x04" + b"A" * (11 * 1024 * 1024)
    resp = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", oversized, "application/octet-stream")}, headers=headers,
    )
    assert resp.status_code == 413


async def test_wrong_file_type_spoofed_extension_rejected_422(client, auth_headers):
    headers = await auth_headers("Engineer")
    resp = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", b"not really an xlsx file", "application/octet-stream")}, headers=headers,
    )
    assert resp.status_code == 422


async def test_row_count_cap_exceeded_rejected(client, auth_headers, monkeypatch):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)

    import app.application.bulk_import.parsing as parsing_module

    monkeypatch.setattr(parsing_module, "MAX_BULK_IMPORT_ROWS", 3)

    rows = [
        [f"RACK-{uuid.uuid4().hex[:8]}", "Rack", model["manufacturer"], model["model_name"], "", room["site_code"],
         room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", ""]
        for _ in range(5)
    ]
    content = build_workbook(RACK_HEADERS, rows)
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "failed_parse"
    assert "maximum allowed" in job_status["rejection_reason"]


async def test_partially_valid_workbook_only_good_rows_commit_report_reflects_both(client, auth_headers):
    headers = await auth_headers("Engineer")
    room = await create_room_with_codes(client, auth_headers)
    model = await _create_rack_model(client, auth_headers)
    good_tag = f"RACK-{uuid.uuid4().hex[:8]}"

    good_row = [good_tag, "Good Rack", model["manufacturer"], model["model_name"], "", room["site_code"],
                room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", ""]
    bad_row = [f"RACK-{uuid.uuid4().hex[:8]}", "Bad Rack", "NoSuchMfr", "NoSuchModel", "", room["site_code"],
               room["building_code"], room["floor_level"], room["room_code"], "", "", "", "", ""]

    content = build_workbook(RACK_HEADERS, [good_row, bad_row])
    upload = await client.post(
        "/api/v1/racks/import-jobs?mode=create_only",
        files={"file": ("racks.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    _run_commit(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed_with_errors"
    assert job_status["committed_row_count"] == 1
    assert job_status["error_row_count"] == 1

    racks = (await client.get("/api/v1/racks", headers=headers)).json()
    assert any(r["asset_tag"] == good_tag for r in racks["items"])

    import io

    from openpyxl import load_workbook

    report_bytes = (await client.get(f"/api/v1/import-jobs/{job['id']}/report", headers=headers)).content
    report_wb = load_workbook(io.BytesIO(report_bytes))
    report_sheet = report_wb["Report"]
    statuses = [row[-5].value for row in report_sheet.iter_rows(min_row=2)]
    assert "committed" in statuses
    assert "invalid" in statuses
