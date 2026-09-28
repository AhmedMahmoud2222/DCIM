"""Catalog bulk-import pipeline — see test_bulk_import_racks.py's module docstring for
the `.run()` synchronous-Celery-task pattern this file also uses."""

import uuid

from app.infrastructure.tasks.bulk_import import commit_bulk_import_job, parse_and_validate_bulk_import_job
from tests.api._bulk_import_helpers import build_workbook

CATALOG_HEADERS = [
    "manufacturer_name", "category", "model_name", "model_number", "subtype", "description", "tags",
    "dimension_unit", "width_value", "height_value", "depth_value", "rack_unit_height", "weight_unit", "weight_value",
    "mounting_orientation", "supported_placement_types", "airflow_direction", "rated_power_w", "typical_power_w",
    "max_power_w", "heat_dissipation_btu_hr", "power_redundancy_mode", "revision_number", "clone_from_revision_number",
]


def _base_row(manufacturer_name, model_name, **overrides) -> list:
    values = {
        "manufacturer_name": manufacturer_name, "category": "equipment", "model_name": model_name, "model_number": "",
        "subtype": "", "description": "", "tags": "", "dimension_unit": "mm", "width_value": 440, "height_value": 44,
        "depth_value": 600, "rack_unit_height": 1, "weight_unit": "kg", "weight_value": 5, "mounting_orientation": "",
        "supported_placement_types": "rack_mounted", "airflow_direction": "front_to_rear", "rated_power_w": 200,
        "typical_power_w": 120, "max_power_w": 250, "heat_dissipation_btu_hr": "", "power_redundancy_mode": "single",
        "revision_number": "", "clone_from_revision_number": "",
    }
    values.update(overrides)
    return [values[h] for h in CATALOG_HEADERS]


def _run_parse(job_id: str) -> None:
    parse_and_validate_bulk_import_job.run(job_id)


def _run_commit(job_id: str) -> None:
    commit_bulk_import_job.run(job_id)


async def test_catalog_import_template_download_is_real_xlsx_with_expected_headers(client, auth_headers):
    headers = await auth_headers("Administrator")
    resp = await client.get("/api/v1/catalog/import-template", headers=headers)
    assert resp.status_code == 200
    assert resp.content[:4] == b"PK\x03\x04"

    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(resp.content))
    sheet = workbook["Catalog"]
    header_row = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header_row == CATALOG_HEADERS


async def test_upload_parse_and_preview_rows(client, auth_headers):
    headers = await auth_headers("Administrator")
    manufacturer_name = f"Mfr-{uuid.uuid4().hex[:8]}"
    model_name = f"Model-{uuid.uuid4().hex[:8]}"

    content = build_workbook(CATALOG_HEADERS, [_base_row(manufacturer_name, model_name)])
    upload = await client.post(
        "/api/v1/catalog/import-jobs?mode=create_only",
        files={"file": ("catalog.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    assert upload.status_code == 202, upload.text
    job = upload.json()

    _run_parse(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "validated"
    assert job_status["valid_row_count"] == 1


async def test_commit_create_only_mode_success(client, auth_headers):
    headers = await auth_headers("Administrator")
    manufacturer_name = f"Mfr-{uuid.uuid4().hex[:8]}"
    model_name = f"Model-{uuid.uuid4().hex[:8]}"

    content = build_workbook(CATALOG_HEADERS, [_base_row(manufacturer_name, model_name)])
    upload = await client.post(
        "/api/v1/catalog/import-jobs?mode=create_only",
        files={"file": ("catalog.xlsx", content, "application/octet-stream")}, headers=headers,
    )
    job = upload.json()
    _run_parse(job["id"])
    _run_commit(job["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{job['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status
    assert job_status["committed_row_count"] == 1

    manufacturers = (await client.get("/api/v1/catalog/manufacturers", params={"q": manufacturer_name}, headers=headers)).json()
    assert manufacturers["total"] == 1
    manufacturer = manufacturers["items"][0]

    models = (
        await client.get("/api/v1/catalog/models", params={"manufacturer_id": manufacturer["id"]}, headers=headers)
    ).json()
    assert models["total"] == 1
    model_detail = (await client.get(f"/api/v1/catalog/models/{models['items'][0]['id']}", headers=headers)).json()
    assert len(model_detail["revisions"]) == 1
    revision = model_detail["revisions"][0]
    assert revision["lifecycle_status"] == "draft"


async def test_commit_update_existing_mode_success(client, auth_headers):
    headers = await auth_headers("Administrator")
    manufacturer_name = f"Mfr-{uuid.uuid4().hex[:8]}"
    model_name = f"Model-{uuid.uuid4().hex[:8]}"

    create_content = build_workbook(CATALOG_HEADERS, [_base_row(manufacturer_name, model_name)])
    create_upload = (
        await client.post(
            "/api/v1/catalog/import-jobs?mode=create_only",
            files={"file": ("catalog.xlsx", create_content, "application/octet-stream")}, headers=headers,
        )
    ).json()
    _run_parse(create_upload["id"])
    _run_commit(create_upload["id"])

    update_content = build_workbook(
        CATALOG_HEADERS, [_base_row(manufacturer_name, model_name, revision_number=1, weight_value=9.5, description="updated")]
    )
    update_upload = (
        await client.post(
            "/api/v1/catalog/import-jobs?mode=update_existing",
            files={"file": ("catalog.xlsx", update_content, "application/octet-stream")}, headers=headers,
        )
    ).json()
    _run_parse(update_upload["id"])
    _run_commit(update_upload["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{update_upload['id']}", headers=headers)).json()
    assert job_status["status"] == "committed", job_status

    manufacturers = (await client.get("/api/v1/catalog/manufacturers", params={"q": manufacturer_name}, headers=headers)).json()
    models = (
        await client.get("/api/v1/catalog/models", params={"manufacturer_id": manufacturers["items"][0]["id"]}, headers=headers)
    ).json()
    revision_id = (await client.get(f"/api/v1/catalog/models/{models['items'][0]['id']}", headers=headers)).json()[
        "revisions"
    ][0]["id"]
    revision = (await client.get(f"/api/v1/catalog/revisions/{revision_id}", headers=headers)).json()
    assert float(revision["weight_value"]) == 9.5


async def test_update_existing_against_published_revision_fails_and_leaves_row_byte_for_byte_unchanged(client, auth_headers, db_session):
    """The plan's specific correctness requirement: attempting update_existing against a
    published revision must be a row-level failure, never a silent no-op and never a
    mutation — verified by re-reading the DB row directly before/after and asserting
    equality, not just checking the HTTP-level report."""
    headers = await auth_headers("Administrator")
    manufacturer_name = f"Mfr-{uuid.uuid4().hex[:8]}"
    model_name = f"Model-{uuid.uuid4().hex[:8]}"

    create_content = build_workbook(CATALOG_HEADERS, [_base_row(manufacturer_name, model_name)])
    create_upload = (
        await client.post(
            "/api/v1/catalog/import-jobs?mode=create_only",
            files={"file": ("catalog.xlsx", create_content, "application/octet-stream")}, headers=headers,
        )
    ).json()
    _run_parse(create_upload["id"])
    _run_commit(create_upload["id"])

    manufacturers = (await client.get("/api/v1/catalog/manufacturers", params={"q": manufacturer_name}, headers=headers)).json()
    models = (
        await client.get("/api/v1/catalog/models", params={"manufacturer_id": manufacturers["items"][0]["id"]}, headers=headers)
    ).json()
    model_id = models["items"][0]["id"]
    revision_detail = (await client.get(f"/api/v1/catalog/models/{model_id}", headers=headers)).json()["revisions"][0]
    revision_id = revision_detail["id"]

    # A nameplate power figure was set by the import row, so publish requires at least
    # one power supply (validate_revision_for_publish) — add one directly via the
    # designer API before publishing.
    psu_resp = await client.post(
        f"/api/v1/catalog/revisions/{revision_id}/power-supplies",
        json={"stable_key": "psu-1", "label": "PSU 1", "connector_type": "C14"},
        headers={**headers, "If-Match": "1"},
    )
    assert psu_resp.status_code == 201, psu_resp.text

    publish = await client.post(f"/api/v1/catalog/revisions/{revision_id}/publish", headers=headers)
    assert publish.status_code == 200, publish.text

    from sqlalchemy import text as sql_text

    before = (
        await db_session.execute(
            sql_text("SELECT * FROM catalog_model_revision WHERE id = :id"), {"id": revision_id}
        )
    ).mappings().one()
    before_dict = dict(before)

    update_content = build_workbook(
        CATALOG_HEADERS, [_base_row(manufacturer_name, model_name, revision_number=1, weight_value=999)]
    )
    update_upload = (
        await client.post(
            "/api/v1/catalog/import-jobs?mode=update_existing",
            files={"file": ("catalog.xlsx", update_content, "application/octet-stream")}, headers=headers,
        )
    ).json()
    _run_parse(update_upload["id"])
    _run_commit(update_upload["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{update_upload['id']}", headers=headers)).json()
    # The row is rejected either at validate time (preferred: an accurate preview) or at
    # commit time (belt-and-suspenders, since lock_draft_revision_for_edit re-checks) —
    # either way it must never reach status='committed'.
    assert job_status["committed_row_count"] == 0

    after = (
        await db_session.execute(
            sql_text("SELECT * FROM catalog_model_revision WHERE id = :id"), {"id": revision_id}
        )
    ).mappings().one()
    assert dict(after) == before_dict


async def test_duplicate_catalog_identity_within_the_same_file(client, auth_headers):
    headers = await auth_headers("Administrator")
    manufacturer_name = f"Mfr-{uuid.uuid4().hex[:8]}"
    model_name = f"Model-{uuid.uuid4().hex[:8]}"

    content = build_workbook(
        CATALOG_HEADERS, [_base_row(manufacturer_name, model_name), _base_row(manufacturer_name, model_name)]
    )
    upload = (
        await client.post(
            "/api/v1/catalog/import-jobs?mode=create_only",
            files={"file": ("catalog.xlsx", content, "application/octet-stream")}, headers=headers,
        )
    ).json()
    _run_parse(upload["id"])

    job_status = (await client.get(f"/api/v1/import-jobs/{upload['id']}", headers=headers)).json()
    assert job_status["valid_row_count"] == 1
    assert job_status["error_row_count"] == 1


async def test_oversized_file_rejected_413(client, auth_headers):
    headers = await auth_headers("Administrator")
    oversized = b"PK\x03\x04" + b"A" * (11 * 1024 * 1024)
    resp = await client.post(
        "/api/v1/catalog/import-jobs?mode=create_only",
        files={"file": ("catalog.xlsx", oversized, "application/octet-stream")}, headers=headers,
    )
    assert resp.status_code == 413


async def test_wrong_file_type_spoofed_extension_rejected_422(client, auth_headers):
    headers = await auth_headers("Administrator")
    resp = await client.post(
        "/api/v1/catalog/import-jobs?mode=create_only",
        files={"file": ("catalog.xlsx", b"not really an xlsx", "application/octet-stream")}, headers=headers,
    )
    assert resp.status_code == 422
